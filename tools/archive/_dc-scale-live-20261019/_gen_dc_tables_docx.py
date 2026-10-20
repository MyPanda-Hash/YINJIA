# -*- coding: utf-8 -*-
"""生成《数采(数据采集)后台表清单》Word 文档(2026-10-19)。"""
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
import datetime

OUT = r"C:\INCER\YINJIA-MES\docs\development\数采后台表清单.docx"
CJK = "微软雅黑"

doc = Document()
# 页边距
for s in doc.sections:
    s.top_margin = Cm(2.2); s.bottom_margin = Cm(2.2)
    s.left_margin = Cm(2.2); s.right_margin = Cm(2.2)

# 正文字体(中文需显式 eastAsia)
normal = doc.styles["Normal"]
normal.font.name = CJK
normal.font.size = Pt(10.5)
normal.element.rPr.rFonts.set(qn("w:eastAsia"), CJK)


def style_heading(level, size, color=(0x1F, 0x3A, 0x5F)):
    st = doc.styles[f"Heading {level}"]
    st.font.name = CJK
    st.font.size = Pt(size)
    st.font.color.rgb = RGBColor(*color)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), CJK)


for lv, sz in ((1, 15), (2, 12.5), (3, 11)):
    style_heading(lv, sz)


def p(text="", bold=False, size=None, space_after=6):
    par = doc.add_paragraph()
    run = par.add_run(text)
    run.bold = bold
    run.font.name = CJK
    run.element.rPr.rFonts.set(qn("w:eastAsia"), CJK)
    if size:
        run.font.size = Pt(size)
    par.paragraph_format.space_after = Pt(space_after)
    return par


def bullet(text):
    par = doc.add_paragraph(text, style="List Bullet")
    for r in par.runs:
        r.font.name = CJK
        r.element.rPr.rFonts.set(qn("w:eastAsia"), CJK)
        r.font.size = Pt(10.5)
    par.paragraph_format.space_after = Pt(2)


def table(headers, rows, widths=None, font=9.5):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    hdr = t.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = ""
        run = hdr[i].paragraphs[0].add_run(h)
        run.bold = True
        run.font.size = Pt(font)
        run.font.name = CJK
        run.element.rPr.rFonts.set(qn("w:eastAsia"), CJK)
    for r in rows:
        cells = t.add_row().cells
        for i, val in enumerate(r):
            cells[i].text = ""
            run = cells[i].paragraphs[0].add_run(str(val))
            run.font.size = Pt(font)
            run.font.name = CJK
            run.element.rPr.rFonts.set(qn("w:eastAsia"), CJK)
    if widths:
        for row in t.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Cm(w)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


# ── 封面标题 ──
title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
tr = title.add_run("数采(数据采集)后台表清单")
tr.bold = True
tr.font.size = Pt(22)
tr.font.name = CJK
tr.element.rPr.rFonts.set(qn("w:eastAsia"), CJK)

sub = doc.add_paragraph()
sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
sr = sub.add_run("YINJIA-MES · 数据抓取落库范围   |   编制日期 " +
                 datetime.date(2026, 10, 19).strftime("%Y-%m-%d"))
sr.font.size = Pt(10)
sr.font.color.rgb = RGBColor(0x60, 0x6A, 0x76)
sr.font.name = CJK
sr.element.rPr.rFonts.set(qn("w:eastAsia"), CJK)

doc.add_paragraph()

# ── 一、说明 ──
doc.add_heading("一、范围与口径", level=1)
p("凡是「硬件采集到的数据」落库的地方都在这 7 张表里，全部使用 bs_dc_ 前缀。"
  "其中台账 4 张（网关 / 设备 / 采集点 / 规则），数据 3 张（采集流水 / 报警记录 / 在线状态日志）。")
p("建表脚本：tools/migrate-dc-platform-20261018.sql；"
  "表前缀与命名规则见 docs/development/数据库规范.md。下表行数为 2026-10-19 15:08 正式账套 HSDZ_MES 的实测快照。")

# ── 二、总表 ──
doc.add_heading("二、表清单总览", level=1)
table(
    ["#", "表名", "中文名", "存什么", "谁写", "保留策略"],
    [
        ["1", "bs_dc_gateway", "数采网关", "协议、IP/端口、在线状态、接入密钥 api_key", "界面 /gateway/save", "台账(4 行)"],
        ["2", "bs_dc_device", "数采设备", "设备台账、绑定网关、在线状态、最近上报时间", "界面 /device/save；上报时刷在线", "台账(4 行)"],
        ["3", "bs_dc_point", "采集点", "采集变量：编码/类型/单位/量程/采集周期", "界面 /point/save", "台账(6 行)"],
        ["4", "bs_dc_rule", "数采规则", "定时采集 / 变化上报 / 阈值报警规则", "界面 /rule/save", "台账(6 行)"],
        ["5", "bs_dc_sample", "采集流水", "每次抓到的值：数值或文本 + 采集时刻", "硬件上报 /api/dc/ingest、/api/dc/ingest/text", "每点留最新 10 条，5 分钟清一次(136 行)"],
        ["6", "bs_dc_alarm", "报警记录", "超限报警：级别、内容、触发值、处理留痕", "上报时自动产生；/alarm/handle", "只增不删(8 行)"],
        ["7", "bs_dc_status_log", "在线状态日志", "在线↔离线状态变更(在线率唯一依据)", "上报/心跳触发状态变化时", "只增不删(2 行)"],
    ],
    widths=[0.8, 3.0, 2.0, 4.6, 3.6, 3.0],
    font=8.5,
)

# ── 三、数据流 ──
doc.add_heading("三、数据是怎么流进来的", level=1)
flow = (
    "电子秤 / 扫码枪 / PLC\n"
    "   │  串口(RS232 9600-8-N-1)或网络\n"
    "   ▼\n"
    "W111 无线透传模块 / 数采网关\n"
    "   │  ① JSON 口  POST /api/dc/ingest         {deviceCode, values:[{pointCode, value|text, at}]}\n"
    "   │  ② 文本口  POST /api/dc/ingest/text    正文 = 串口帧原文(如 +000.41000.00kg)\n"
    "   │  ③ 心跳    POST /api/dc/heartbeat      {deviceCode}\n"
    "   ▼\n"
    "MES(DcController → DcService)\n"
    "   ├─ 写 bs_dc_sample            一条采集流水(值 + 采集时刻)\n"
    "   ├─ 刷 bs_dc_device            在线状态 / 最近上报 / 最近调试\n"
    "   ├─ 状态变化时写 bs_dc_status_log\n"
    "   └─ 越过采集点上下限时写 bs_dc_alarm(预警、未处理)"
)
fp = doc.add_paragraph()
fr = fp.add_run(flow)
fr.font.name = "Consolas"
fr.font.size = Pt(8.5)
fr.element.rPr.rFonts.set(qn("w:eastAsia"), CJK)
fp.paragraph_format.left_indent = Cm(0.4)
fp.paragraph_format.space_after = Pt(8)

p("文本上报口的两道闸(2026-10-19 用户口径)：", bold=True)
bullet("空重量(≤5g)不写流水，但刷新一次在线心跳，避免空秤被判离线；")
bullet("同一设备 1 秒最多写一条，其余帧丢弃(秤连续发送约 2 帧/秒)。")

# ── 四、逐表字段 ──
doc.add_heading("四、逐表关键字段", level=1)

doc.add_heading("4.1 bs_dc_sample(数据抓取主表)", level=2)
table(
    ["列", "类型", "说明"],
    [
        ["id", "bigint 自增", "主键；流水只增不改"],
        ["device_id / point_id", "int", "哪台设备的哪个采集点"],
        ["value_num", "decimal(18,4)", "数值型采集值(重量、读码率等)"],
        ["value_text", "nvarchar(200)", "文本型采集值(条码内容等)"],
        ["quality", "char(1)", "质量位，默认 G"],
        ["collect_at", "datetime2", "采集时刻(不是入库时刻；回放历史时用 at 指定)"],
        ["asp_user1", "nvarchar(50)", "来源标记(ingest)"],
    ],
    widths=[4.2, 3.6, 9.2],
)
p("索引：PK_bs_dc_sample(id)、IX_bs_dc_sample_point_at(point_id, collect_at DESC)。")

doc.add_heading("4.2 其他数据表", level=2)
table(
    ["表", "关键列", "说明"],
    [
        ["bs_dc_alarm", "level / message / value_text / occur_at / state / handled_by / handled_at / handled_note", "报警工单：发生后待处理，处理时写处理人与说明"],
        ["bs_dc_status_log", "device_id / status / changed_at / reason", "只在状态变化时写；状态统计的在线率与在线时长由此折叠"],
    ],
    widths=[3.2, 8.0, 5.8],
)

doc.add_heading("4.3 台账表关键列", level=2)
table(
    ["表", "关键列", "说明"],
    [
        ["bs_dc_gateway", "code / protocol / host / port / status / api_key", "api_key 是硬件上报的唯一凭据：请求头 X-DC-Key 必须匹配"],
        ["bs_dc_device", "code / device_type / gateway_id / status / last_seen_at / heartbeat_seconds", "离线判定 = 最近上报早于 heartbeat_seconds×3"],
        ["bs_dc_point", "device_id / code / data_type / unit / freq_seconds / low_limit / high_limit", "上报用 code；越限才写报警"],
        ["bs_dc_rule", "device_id / name / rule_type / expression / params", "规则配置，供后续定时采集/报警扩展"],
    ],
    widths=[3.2, 8.0, 5.8],
)

# ── 五、保留策略 ──
doc.add_heading("五、保留与清理策略", level=1)
table(
    ["表", "保留策略", "配置项"],
    [
        ["bs_dc_sample", "每个采集点只保留最新 10 条，每 5 分钟清理一次(DcSampleCleanupScheduler，cron 0 */5 * * * *)", "yinjia.dc.sample.retain(默认 10，≤0 关闭)；yinjia.dc.sample.cleanup-cron"],
        ["其余 6 张", "不自动清理：台账 + 报警 + 状态日志都是业务凭据", "—"],
    ],
    widths=[3.2, 8.4, 5.4],
)
p("清理只删流水表，不动报警、状态日志与台账，因此状态统计/在线率不受影响。")

# ── 六、当前快照 ──
doc.add_heading("六、当前快照(2026-10-19 15:08)", level=1)
table(
    ["设备", "状态", "采集点", "说明"],
    [
        ["DZ-001 电子秤", "online", "W01 重量(kg, 2s)；P01 为早期遗留空点", "天翊 DZ-06 计价秤，经 W111 HTTP 直推 MES"],
        ["FX-001 扫码枪", "online", "SCAN_CODE / SCAN_RATE / SCAN_COUNT", "验收演示数据，可删"],
        ["SIM-SCALE-01 测试电子秤", "online(旧数据)", "W01 重量", "上一轮联调遗留"],
        ["KB-001 车间看板", "offline", "—", "台账占位，尚未接入"],
    ],
    widths=[3.6, 2.4, 5.4, 5.6],
)
p("表行数：gateway 4 / device 4 / point 6 / rule 6 / sample 136(持续 1 秒一条、5 分钟一清)/ alarm 8 / status_log 2。")

# ── 七、相关文件 ──
doc.add_heading("七、相关文档与脚本", level=1)
bullet("接入说明：docs/development/数采接入-W111模块.md(模块配置、帧格式、常用命令)")
bullet("建表脚本：tools/migrate-dc-platform-20261018.sql(7 张表 + 中文注明 + 首批设备台账)")
bullet("全库表清单：docs/development/数据库表清单.md；在册台账：tools/db-inuse-tables.txt")
bullet("取数与取证脚本：tools/archive/_dc-scale-live-20261019/(TCP 抓包、HTTP 接收端、模块配置、活体截图)")

doc.save(OUT)

# 自检
d = Document(OUT)
paras = len(d.paragraphs)
tables = len(d.tables)
text = "\n".join(p.text for p in d.paragraphs)
assert "数采(数据采集)后台表清单" in text
assert "bs_dc_sample" in text
assert tables == 6, f"表格数异常: {tables}"
print(f"OK -> {OUT}\n段落 {paras} 个, 表格 {tables} 个")
