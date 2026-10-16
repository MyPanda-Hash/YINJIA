"""读取《生产相关品质检验.xlsx》的结构:sheet 清单、每表维度、非空区域。"""
import openpyxl
import json
import sys

SRC = r"C:\Users\vigna\Documents\WeChat Files\wxid_9he2bhyytczm22\FileStorage\File\2026-10\生产相关品质检验.xlsx"

wb = openpyxl.load_workbook(SRC, data_only=True)
print("SHEETS=" + json.dumps(wb.sheetnames, ensure_ascii=False))
for ws in wb.worksheets:
    print("=" * 70)
    print(f"SHEET: {ws.title}  dims={ws.dimensions}  max_row={ws.max_row} max_col={ws.max_column}")
    print(f"  merged={len(ws.merged_cells.ranges)}")
    # 统计非空单元格
    n = 0
    for row in ws.iter_rows():
        for c in row:
            if c.value is not None and str(c.value).strip() != "":
                n += 1
    print(f"  nonempty_cells={n}")
