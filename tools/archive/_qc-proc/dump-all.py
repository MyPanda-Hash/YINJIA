"""完整 dump 《生产相关品质检验.xlsx》每个 sheet 的非空单元格与合并区域。"""
import openpyxl
import json

SRC = r"C:\Users\vigna\Documents\WeChat Files\wxid_9he2bhyytczm22\FileStorage\File\2026-10\生产相关品质检验.xlsx"

wb = openpyxl.load_workbook(SRC, data_only=True)
out = []
for ws in wb.worksheets:
    out.append("=" * 78)
    out.append(f"### SHEET: {ws.title}   ({ws.max_row} 行 x {ws.max_column} 列)")
    out.append("-- 合并区域 --")
    for mr in sorted(ws.merged_cells.ranges, key=lambda r: (r.min_row, r.min_col)):
        out.append(f"   {mr}")
    out.append("-- 单元格(按行) --")
    for row in ws.iter_rows():
        cells = []
        for c in row:
            v = c.value
            if v is None:
                continue
            s = str(v).replace("\n", "\\n").strip()
            if s == "":
                continue
            cells.append(f"{c.coordinate}={s}")
        if cells:
            out.append("  " + " | ".join(cells))

txt = "\n".join(out)
with open(r"D:\workspace\yinjia\tools\archive\_qc-proc\dump.txt", "w", encoding="utf-8") as f:
    f.write(txt)
print("WROTE lines=" + str(len(out)))
print(txt)
