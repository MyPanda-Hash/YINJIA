import openpyxl, sys, io

PATH = r'C:\INCER\YINJIA-MES\_scratch\req.xlsx'
OUT = r'C:\INCER\YINJIA-MES\_scratch\req-dump.txt'

wb = openpyxl.load_workbook(PATH, data_only=True)
buf = io.StringIO()
for ws in wb.worksheets:
    buf.write(f'\n{"="*100}\n===== SHEET: {ws.title}   ({ws.max_row} 行 x {ws.max_column} 列) =====\n{"="*100}\n')
    mrs = [str(m) for m in ws.merged_cells.ranges]
    if mrs:
        buf.write(f'[合并区 {len(mrs)} 个] ' + ', '.join(mrs) + '\n\n')
    for r in ws.iter_rows():
        cells = [(c.coordinate, c.value) for c in r if c.value not in (None, '')]
        if not cells:
            continue
        buf.write(f'--- 行{r[0].row} ---\n')
        for k, v in cells:
            s = str(v).replace('\r\n', ' / ').replace('\n', ' / ')
            buf.write(f'  {k}: {s}\n')

with open(OUT, 'w', encoding='utf-8') as f:
    f.write(buf.getvalue())
print('written', OUT, len(buf.getvalue()), 'chars')
