# _prod-forms-extract.py — 提取生产单据 Excel 的文字与嵌入图片
import sys, io, os, zipfile, re
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8")
import openpyxl

ROOT = r"C:\Users\x1787\Documents\xwechat_files\wxid_i0jnge0qyt5w22_b496\msg\file\2026-09\生产相关表格和生产排单表"
OUT = r"C:\INCER\YINJIA-MES\tools\_flow-v12\forms-images"
os.makedirs(OUT, exist_ok=True)

files = {}
for dirpath, _d, names in os.walk(ROOT):
    for n in names:
        if n.lower().endswith(".xlsx"):
            files[n] = os.path.join(dirpath, n)

# ── 1. Sheet1 文字内容(生产相关表格文件.xlsx) ──
big = files["生产相关表格文件.xlsx"]
wb = openpyxl.load_workbook(big, read_only=True)
ws = wb["Sheet1"]
print("===== 生产相关表格文件.xlsx / Sheet1 全部文字 =====")
for i, row in enumerate(ws.iter_rows(values_only=True), 1):
    cells = [str(c).replace("\n", "\\n") for c in row if c is not None and str(c).strip()]
    if cells:
        print(f"  r{i}: {' | '.join(cells)}")
wb.close()

# ── 2. 排单计划表格.xlsx / 成型排单计划 表头与前几行 ──
pd = files["排单计划表格.xlsx"]
wb = openpyxl.load_workbook(pd, read_only=True)
ws = wb["成型排单计划"]
print("\n===== 排单计划表格.xlsx / 成型排单计划(前20行,每行最多30个非空) =====")
for i, row in enumerate(ws.iter_rows(max_row=20, values_only=True), 1):
    cells = [str(c).replace("\n", "\\n")[:22] for c in row if c is not None and str(c).strip()]
    if cells:
        print(f"  r{i}({len(cells)}): {' | '.join(cells[:30])}")
wb.close()

# ── 3. 解包嵌入图片 + 锚点行号 ──
print("\n===== 嵌入图片提取 =====")
z = zipfile.ZipFile(big)
media = [n for n in z.namelist() if n.startswith("xl/media/") and not n.endswith("/")]
print("media files:", len(media))
# drawing 锚点: 从 xl/drawings/drawing1.xml 读 row 与 rId 关系
anchors = {}
for dn in [n for n in z.namelist() if re.match(r"xl/drawings/drawing\d+\.xml", n)]:
    xml = z.read(dn).decode("utf-8", "ignore")
    rel_path = dn.replace("drawings/", "drawings/_rels/") + ".rels"
    rid2file = {}
    if rel_path in z.namelist():
        rels = z.read(rel_path).decode("utf-8", "ignore")
        for m in re.finditer(r'Id="(rId\d+)"[^>]*Target="\.\./media/([^"]+)"', rels):
            rid2file[m.group(1)] = m.group(2)
    for m in re.finditer(r"<xdr:twoCellAnchor.*?</xdr:twoCellAnchor>", xml, re.S):
        blk = m.group(0)
        row_m = re.search(r"<xdr:from>.*?<xdr:row>(\d+)</xdr:row>", blk, re.S)
        rid_m = re.search(r'r:embed="(rId\d+)"', blk)
        if row_m and rid_m and rid_m.group(1) in rid2file:
            anchors[rid2file[rid_m.group(1)]] = int(row_m.group(1)) + 1

for mf in sorted(media):
    fname = os.path.basename(mf)
    row_no = anchors.get(fname, "?")
    dest = os.path.join(OUT, f"r{row_no}_{fname}" if row_no != "?" else fname)
    with open(dest, "wb") as f:
        f.write(z.read(mf))
print("extracted:", len(media), "-> tools/_flow-v12/forms-images/ (r{锚点行}_{原名})")
z.close()
