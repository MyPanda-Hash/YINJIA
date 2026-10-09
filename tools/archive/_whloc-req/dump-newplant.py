import openpyxl, io, json

XLSX = r'C:\Users\vigna\Documents\WeChat Files\wxid_9he2bhyytczm22\FileStorage\File\2026-09\仓库库位信息表.xlsx'
wb = openpyxl.load_workbook(XLSX, data_only=True)
ws = wb['Sheet2']

merged = {}
for rng in ws.merged_cells.ranges:
    for r in range(rng.min_row, rng.max_row + 1):
        for c in range(rng.min_col, rng.max_col + 1):
            merged[(r, c)] = str(rng)

buf = io.StringIO()
buf.write('=== 新厂区 (行 30-48) 逐格 dump;每格: 值 [合并区] ===\n')
for r in range(30, 49):
    cells = []
    for c in range(1, 16):
        v = ws.cell(row=r, column=c).value
        s = '' if v is None else str(v).strip()
        m = merged.get((r, c))
        if not s and not m:
            continue
        cells.append(f'{openpyxl.utils.get_column_letter(c)}{r}={s!r}' + (f'[{m}]' if m else ''))
    if cells:
        buf.write('  ' + ' | '.join(cells) + '\n')

buf.write('\n=== 新厂区 表头行 31/32 与数据列对照 ===\n')
for c in range(1, 16):
    L = openpyxl.utils.get_column_letter(c)
    def g(r):
        v = ws.cell(row=r, column=c).value
        return '' if v is None else str(v).strip()
    buf.write(f'  列{L}: r31={g(31)!r}  r32={g(32)!r}\n')

buf.write('\n=== 各数据列的非空单元格(行33-48) ===\n')
for c in range(1, 16):
    L = openpyxl.utils.get_column_letter(c)
    vals = []
    for r in range(33, 49):
        v = ws.cell(row=r, column=c).value
        s = '' if v is None else str(v).strip()
        if s:
            vals.append(f'r{r}:{s}')
    if vals:
        buf.write(f'  列{L}: ' + ' ; '.join(vals) + '\n')

open(r'D:\workspace\yinjia\tools\archive\_whloc-req\dump-newplant.txt', 'w', encoding='utf-8').write(buf.getvalue())
print('ok')
