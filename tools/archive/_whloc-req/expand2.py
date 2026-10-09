import openpyxl, re, io, collections

p = r'C:\Users\vigna\Documents\WeChat Files\wxid_9he2bhyytczm22\FileStorage\File\2026-09\仓库库位信息表.xlsx'
wb = openpyxl.load_workbook(p, data_only=True)
ws = wb['Sheet2']

def val(r, c):
    v = ws.cell(row=r, column=c).value
    return '' if v is None else str(v).strip()

fill = {}
for rng in ws.merged_cells.ranges:
    v = ws.cell(row=rng.min_row, column=rng.min_col).value
    for r in range(rng.min_row, rng.max_row + 1):
        for c in range(rng.min_col, rng.max_col + 1):
            fill[(r, c)] = '' if v is None else str(v).strip()
def g(r, c):
    return fill.get((r, c), val(r, c))

HDR = {c: g(5, c) for c in range(1, 16)}
HDR2 = {c: g(32, c) for c in range(1, 16)}

def cells(r0, r1):
    out = collections.defaultdict(list)
    for c in range(1, 16):
        for r in range(r0, r1 + 1):
            t = val(r, c)
            if t:
                out[c].append((r, t))
    return out

NONSTORE = ('柱子', '设备区', '成品区', '呆滞区', '不良品', '过道', '洗手间', '办公室')

def expand_old(t):
    t = re.sub(r'\s+', '', t.replace('\n', ''))
    if not t: return []
    if t in NONSTORE: return []
    if '库位排列一样' in t or t.endswith('...'): return None  # 待补
    parts = t.split('-')
    if len(parts) == 3:
        head, mid, seg = parts
        if '/' in seg: return ['%s-%s-%s' % (head, mid, x) for x in seg.split('/') if x]
        return ['%s-%s-%s' % (head, mid, seg)]
    if len(parts) == 2:
        head, seg = parts
        if '/' in seg: return ['%s-%s' % (head, x) for x in seg.split('/') if x]
        return ['%s-%s' % (head, seg)]
    return []

def expand_new(t):
    t = re.sub(r'\s+', '', t.replace('\n', ''))
    if not t: return [], None
    if t in NONSTORE: return [], 'nonstore'
    m = re.fullmatch(r'D(\d+)[-~]D(\d+)库位排列一样', t)
    if m:
        return [], ('inherit', int(m.group(1)), int(m.group(2)))
    # 区前缀 + 区号 + 区间: 炭1区1-36 / 胶粉2区1-36 / 1-36辅料 / 1-14 / 1-7 / 1-12 / 1-24
    m = re.fullmatch(r'.*?(\d+)[-~](\d+)(辅料)?', t)
    if m and not re.match(r'^[A-Za-z]', t):
        a, b = int(m.group(1)), int(m.group(2))
        if 0 < b - a < 100:
            return ['%s#%d' % (t, i) for i in range(a, b + 1)], 'range'
    # 货架区 H1-1/2/3/4 (H1..H11)
    m = re.fullmatch(r'H(\d+)-(.+)', t)
    if m:
        seg = m.group(2)
        return ['H%s-%s' % (m.group(1), x) for x in seg.split('/') if x], 'ok'
    # D1-01-1/2
    parts = t.split('-')
    if len(parts) == 3 and '/' in parts[2]:
        return ['%s-%s-%s' % (parts[0], parts[1], x) for x in parts[2].split('/') if x], 'ok'
    if len(parts) == 2 and '/' in parts[1]:
        return ['%s-%s' % (parts[0], x) for x in parts[1].split('/') if x], 'ok'
    return [t], 'verbatim'

buf = io.StringIO()

# ---------- 老厂区 ----------
old = cells(6, 29)
old_atoms, old_pending = set(), []
per_col = {}
for c in sorted(old):
    atoms = []
    for r, t in old[c]:
        e = expand_old(t)
        if e is None:
            old_pending.append((openpyxl.utils.get_column_letter(c), HDR.get(c, ''), t))
            continue
        atoms.extend(e)
    per_col[c] = (HDR.get(c, ''), len(set(atoms)))
    old_atoms |= set(atoms)

buf.write('===== 老厂区(Sheet2 R6-R29,列A-O) =====\n')
for c in sorted(per_col):
    if per_col[c][1]:
        buf.write('  列%s %-14s %d\n' % (openpyxl.utils.get_column_letter(c), per_col[c][0], per_col[c][1]))
buf.write('  老厂区已明确库位 = %d; 待业务补全单元格 = %d\n' % (len(old_atoms), len(old_pending)))
for col, hdr, t in old_pending:
    buf.write('     - 列%s %-10s %s\n' % (col, hdr, t))

# ---------- 新厂区 ----------
new = cells(33, 48)
new_atoms = set()
inherit = []
per_col2 = collections.OrderedDict()
for c in sorted(new):
    atoms = []
    for r, t in new[c]:
        e, kind = expand_new(t)
        if isinstance(kind, tuple) and kind[0] == 'inherit':
            inherit.append((c, HDR2.get(c, ''), kind[1], kind[2]))
            continue
        atoms.extend(e)
    per_col2[c] = (HDR2.get(c, ''), len(set(atoms)))
    new_atoms |= set(atoms)

buf.write('\n===== 新厂区 D仓(Sheet2 R33-R48,列A-O) =====\n')
for c in sorted(per_col2):
    if per_col2[c][1]:
        buf.write('  列%s %-12s %d\n' % (openpyxl.utils.get_column_letter(c), per_col2[c][0], per_col2[c][1]))
buf.write('  新厂区已明确库位 = %d\n' % len(new_atoms))
buf.write('  继承规则 %d 条(需按模板复制):\n' % len(inherit))
for c, hdr, a, b in inherit:
    buf.write('     - 列%s %-10s D%d..D%d 与左侧模板同排列\n' % (openpyxl.utils.get_column_letter(c), hdr, a, b))

buf.write('\n===== 合计 =====\n')
buf.write('已明确库位 = %d(老 %d + 新 %d)\n' % (len(old_atoms) + len(new_atoms), len(old_atoms), len(new_atoms)))

open(r'D:\workspace\yinjia\tools\archive\_whloc-req\expand2.txt', 'w', encoding='utf-8').write(buf.getvalue())
print(buf.getvalue()[:400])
