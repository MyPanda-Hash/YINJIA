"""从《仓库库位信息表》Sheet2 生成剩余仓位的 SQL(老厂区成品B/C仓 + 辅料及配件区(A仓) + 新厂区成品D仓)。

- 数据源 = xlsx 原文(非截图),编码严格照原文,不另发明
- 两处原文笔误按正字收录:`B-1-18-1/2/3`⇒B1-18-*、`D2-D7库位排列一样`按 D2/D3 的排列复制
- AH 系列截断(`AH1-1-1/2...`)按用户已确认口径取 **3 层**
- 输出:INSERT 语句(显式行,可审计)+ 分段注释
"""
import openpyxl, re, io, collections

XLSX = r'C:\Users\vigna\Documents\WeChat Files\wxid_9he2bhyytczm22\FileStorage\File\2026-09\仓库库位信息表.xlsx'
wb = openpyxl.load_workbook(XLSX, data_only=True)
ws = wb['Sheet2']

fill = {}
for rng in ws.merged_cells.ranges:
    v = ws.cell(row=rng.min_row, column=rng.min_col).value
    for r in range(rng.min_row, rng.max_row + 1):
        for c in range(rng.min_col, rng.max_col + 1):
            fill[(r, c)] = '' if v is None else str(v).strip()

def g(r, c):
    v = fill.get((r, c), ws.cell(row=r, column=c).value)
    return '' if v is None else str(v).strip()

LAYERS = 3           # AH 截断项的层数(用户口径)
TYPO = {'B-1-18-1/2/3': 'B1-18-1/2/3'}   # 原文笔误

# 老厂区列定义:(列号, 仓库编码, 仓库名称, 存储分区, 区码)
OLD_COLS = {
    1:  ('CP-02', '成品B仓', '成品区', 'B1'),
    2:  ('CP-02', '成品B仓', '成品区', 'B2'),
    3:  ('CP-02', '成品B仓', '成品区', 'B3'),
    4:  ('CP-02', '成品B仓', '成品区', 'B4'),
    5:  ('CK-C',  '成品C仓', '成品区', 'C1'),
    6:  ('CK-C',  '成品C仓', '成品区', 'C2'),
    # 7/8/9 = 炭粉区/胶粉区/货架区(已落库,跳过)
    10: ('CK-A2', '辅料及配件区A仓', '纸箱区',   'A1'),
    11: ('CK-A2', '辅料及配件区A仓', '端盖区',   'A2'),
    12: ('CK-A2', '辅料及配件区A仓', 'PP棉区',   None),   # 原文空白
    13: ('CK-A2', '辅料及配件区A仓', '无纺布区', 'AH'),
    14: ('CK-A2', '辅料及配件区A仓', '网套区',   'AH'),
    15: ('CK-A2', '辅料及配件区A仓', '标签区',   'AH'),
}
NEW_COLS = {   # 新厂区成品仓(D仓);1=D1 2=D2 3=D3 (4=继承说明) 5=D8 (6=继承说明)
    1: ('CK-D', '成品D仓', '成品区', 'D1'),
    2: ('CK-D', '成品D仓', '成品区', 'D2'),
    3: ('CK-D', '成品D仓', '成品区', 'D3'),
    5: ('CK-D', '成品D仓', '成品区', 'D8'),
}

def lead_letters(s):
    m = re.match(r'^[A-Za-z]+', s)
    return len(m.group(0)) if m else 0

def expand(token):
    """把原文 token 展开成 [(区码, 排号, 位号, 层号, 仓位编码)]"""
    t = re.sub(r'\s+', '', token.replace('\n', ''))
    t = TYPO.get(t, t)
    if not t:
        return []
    if t in ('柱子', '设备区', '成品区', '呆滞区', '不良品', '过道', '洗手间', '办公室'):
        return []
    if '库位排列一样' in t:
        return []                                   # 继承说明,单独处理
    trunc = t.endswith('...')
    t = t.replace('...', '')
    parts = t.split('-')
    if len(parts) < 2:
        return []

    def fam(code_head):
        return 'rack' if lead_letters(code_head) >= 2 else 'area'

    out = []
    if len(parts) == 3:
        head, mid, seg = parts
        fam_ = fam(head)
        if fam_ == 'rack':          # <区码><排>-<位>-<层>  例如 AH5-1/2/3-1... 或 AH1-1-1/2...
            zone = re.match(r'^[A-Za-z]+', head).group(0)
            row = head[len(zone):]
            bits = [x for x in mid.split('/') if x]
            lays = [x for x in seg.split('/') if x]
            if trunc:               # 层被截断 ⇒ 取 1..LAYERS
                lays = [str(i) for i in range(1, LAYERS + 1)]
            for b in bits:
                for ly in lays:
                    out.append((zone, row, b, ly, f'{zone}{row}-{b}-{ly}'))
        else:                       # <区码>-<排>-<位>  例如 A1-1/2/3-… 不存在;此处为 B1-01-1/2 形态
            zone = head
            bits = [x for x in seg.split('/') if x]
            for b in bits:
                out.append((zone, mid, b, None, f'{zone}-{mid}-{b}'))
        return out

    if len(parts) == 2:
        head, seg = parts
        fam_ = fam(head)
        if fam_ == 'rack':          # H1-1/2/3/4 (新厂区货架,本次不处理)
            zone = re.match(r'^[A-Za-z]+', head).group(0)
            row = head[len(zone):]
            for b in [x for x in seg.split('/') if x]:
                out.append((zone, row, b, None, f'{zone}{row}-{b}'))
        else:                       # <区码>-<位…> 的简写(本次不出现)
            for b in [x for x in seg.split('/') if x]:
                out.append((head, seg, b, None, f'{head}-{seg}-{b}'))
        return out
    return out

rows = []      # (仓库编码, 仓库名称, 厂区, 存储分区, 区码, 排号, 位号, 层号, 仓位编码)
def emit(code, wh, whname, area, c, r, b, ly):
    rows.append((wh, whname, '老厂区' if wh.startswith(('CP-', 'CK-A', 'CK-C')) else '新厂区',
                 area, c, r, b, ly, code))

# ── 老厂区 R6..R29 ──
old_pending = []
for col, (wh, whname, area, zone) in OLD_COLS.items():
    for r in range(6, 30):
        tok = g(r, col)
        if not tok:
            continue
        if '库位排列一样' in tok:
            old_pending.append((col, tok)); continue
        for (z, row, b, ly, code) in expand(tok):
            emit(code, wh, whname, area, z, row, b, ly)

# ── 新厂区 R33..R48(成品仓列;继承说明单独展开) ──
new_atoms = {}
for col, (wh, whname, area, zone) in NEW_COLS.items():
    for r in range(33, 49):
        tok = g(r, col)
        if not tok:
            continue
        if '库位排列一样' in tok:
            continue
        got = expand(tok)
        new_atoms.setdefault(col, []).extend(got)
        for (z, row, b, ly, code) in got:
            emit(code, wh, whname, area, z, row, b, ly)

# 继承:D2-D7 排列一样(D2/D3 已是 6排×2位) / D8-D11 排列一样(D8 已是 8排×2位)
def inherit(template_cols, targets):
    tpl = []
    for tc in template_cols:
        tpl.extend(new_atoms.get(tc, []))
    seen, base = set(), []
    for a in tpl:
        if a[4] not in seen:
            seen.add(a[4]); base.append(a)
    for tz in targets:
        for (z, row, b, ly, code) in base:
            emit(f'{tz}-{row}-{b}', 'CK-D', '成品D仓', '成品区', tz, row, b, ly)
inherit([2, 3], ['D4', 'D5', 'D6', 'D7'])
inherit([5], ['D9', 'D10', 'D11'])

# 去重(同编码只留一条)
uniq = {}
for r in rows:
    uniq.setdefault(r[8], r)
rows = list(uniq.values())
rows.sort(key=lambda x: (x[0], x[3], x[4], x[5], x[6], x[7] or ''))

buf = io.StringIO()
buf.write('-- 由 tools/archive/_whloc-req/gen-rest.py 自动生成(数据源=《仓库库位信息表》Sheet2 原文)\n')
buf.write(f'-- 合计 {len(rows)} 个仓位\n')

by = collections.Counter((r[0], r[3]) for r in rows)
buf.write('-- 分仓分区计数:\n')
for (wh, area), n in sorted(by.items()):
    buf.write(f'--   {wh:8} {area:8} {n:4}\n')

buf.write("""
DECLARE @loc TABLE (
  仓库编码 nvarchar(100), 仓库名称 nvarchar(200), 厂区 nvarchar(20), 存储分区 nvarchar(30),
  区码 nvarchar(20), 排号 nvarchar(10), 位号 nvarchar(10), 层号 nvarchar(10), 仓位编码 nvarchar(100));

INSERT INTO @loc (仓库编码, 仓库名称, 厂区, 存储分区, 区码, 排号, 位号, 层号, 仓位编码) VALUES
""")
vals = []
for (wh, whn, fac, area, c, r_, b, ly, code) in rows:
    lyv = 'NULL' if ly in (None, '') else f"N'{ly}'"
    vals.append(f"  (N'{wh}',N'{whn}',N'{fac}',N'{area}',N'{c}',N'{r_}',N'{b}',{lyv},N'{code}')")
buf.write(',\n'.join(vals) + ';\n')

buf.write("""
DECLARE @ins int;
INSERT INTO dbo.bs_wh_loc (仓库, 仓库编码, 仓位编码, 仓位地址, 厂区, 区码, 存储分区, 排号, 位号, 层号,
                           停用, asp_user1, asp_time1, asp_cancel)
SELECT l.仓库名称, l.仓库编码, l.仓位编码,
       l.存储分区 + l.区码 + N'-' + l.排号 + N'排' + l.位号 + N'位'
         + ISNULL(l.层号 + N'层', N'')                  -- 地址=存储分区+区码-排+排+位+位[+层+层];无层则不加
         ,
       l.厂区, l.区码, l.存储分区, l.排号, l.位号, l.层号, 0, N'migration', SYSDATETIME(), N'N'
FROM @loc l
WHERE NOT EXISTS (SELECT 1 FROM dbo.bs_wh_loc x
                  WHERE x.仓位编码 = l.仓位编码 AND ISNULL(x.asp_cancel,'N') <> 'Y');
SET @ins = @@ROWCOUNT;
PRINT N'  ✓ 新增仓位 ' + CAST(@ins AS nvarchar(10)) + N' 行';
""")

out = r'D:\workspace\yinjia\tools\archive\_whloc-req\gen-rest-rows.sql'
open(out, 'w', encoding='utf-8').write(buf.getvalue())
print('rows =', len(rows))
for (wh, area), n in sorted(by.items()):
    print(f'  {wh:8} {area:8} {n:4}')
print('written ->', out)
