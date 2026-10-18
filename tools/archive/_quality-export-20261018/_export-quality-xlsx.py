#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""
_export-quality-xlsx.py — 「我的桌面 → 质量」数据导出为 Excel 并随仓库推送到云端(2026-10-18)

用户口径:「我的桌面的质量模块的数据上传GitHub」。
沿用 2026-10-18 职员档案的先例(db/emp-data-20261018/):导出 → 放 db/<主题>-<日期>/ → 提交推送,
供同事拉取后直接打开,不必连库。

数据源:后端 `GET /api/dashboard/stats?from&to` 返回的 `quality` 段(就是「我的桌面」质量卡
与「质量」页签渲染的那份数,同源同口径,不另算一遍):
  total     送检数量
  pass      合格数量
  bad       非合格数量
  passRate  合格率(%)
  byResult  检验结果分布(合格 / 不合格)
  defectItems 不合格物料 TOP(物料/供应商名 + 数量)
  trend7    逐日「送检 / 合格」

用法(仓库根;需 8090 在跑 + 管理员账号):
    $env:YINJIA_MES_USER='admin'; $env:YINJIA_MES_PASS='***'
    python tools/archive/_quality-export-20261018/_export-quality-xlsx.py
可选:--base http://127.0.0.1:8090  --from 2026-09-19 --to 2026-10-18  --out db/quality-data-20261018
"""
import argparse
import datetime as dt
import json
import os
import sys
import urllib.request

ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), '..', '..', '..'))


def http_json(url, payload=None, token=None):
    data = None
    headers = {'Content-Type': 'application/json; charset=utf-8'}
    if payload is not None:
        data = json.dumps(payload).encode('utf-8')
    if token:
        headers['Authorization'] = 'Bearer ' + token
    req = urllib.request.Request(url, data=data, headers=headers)
    with urllib.request.urlopen(req, timeout=30) as r:
        return json.loads(r.read().decode('utf-8'))


def login(base, user, pwd):
    j = http_json(base + '/api/auth/login', {'userName': user, 'password': pwd, 'factory': 'YJ'})
    if j.get('code') != 200 or not j.get('data', {}).get('token'):
        raise SystemExit('[FAIL] 登录失败: ' + json.dumps(j, ensure_ascii=False)[:300])
    return j['data']['token']


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--base', default='http://127.0.0.1:8090')
    ap.add_argument('--from', dest='frm', default=None)
    ap.add_argument('--to', dest='to', default=None)
    ap.add_argument('--out', default=None)
    args = ap.parse_args()

    user = os.environ.get('YINJIA_MES_USER', '')
    pwd = os.environ.get('YINJIA_MES_PASS', '')
    if not user or not pwd:
        print('[FAIL] 请先设环境变量 YINJIA_MES_USER / YINJIA_MES_PASS(不把口令写进脚本/仓库)')
        return 1

    to = args.to or dt.date.today().isoformat()
    frm = args.frm or (dt.date.fromisoformat(to) - dt.timedelta(days=29)).isoformat()
    outdir = args.out or os.path.join(ROOT, 'db', 'quality-data-' + to.replace('-', ''))
    os.makedirs(outdir, exist_ok=True)

    token = login(args.base, user, pwd)
    j = http_json('%s/api/dashboard/stats?from=%s&to=%s' % (args.base, frm, to), token=token)
    if j.get('code') != 200:
        print('[FAIL] 取桌面数据失败: ' + json.dumps(j, ensure_ascii=False)[:300])
        return 1
    d = j['data']
    q = d.get('quality') or {}
    rng = d.get('range') or {}
    print('区间: %s ~ %s (bucket=%s, %s 天)' % (rng.get('from'), rng.get('to'), rng.get('bucket'), rng.get('days')))
    print('质量: 送检=%s 合格=%s 非合格=%s 合格率=%s%%' % (q.get('total'), q.get('pass'), q.get('bad'), q.get('passRate')))

    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = openpyxl.Workbook()
    HEAD_FILL = PatternFill('solid', fgColor='1F4E3D')
    HEAD_FONT = Font(color='FFFFFF', bold=True, size=10)

    def sheet(title, headers, rows, widths=None, freeze='A2'):
        ws = wb.create_sheet(title)
        ws.append(headers)
        for c in range(1, len(headers) + 1):
            cell = ws.cell(row=1, column=c)
            cell.fill = HEAD_FILL
            cell.font = HEAD_FONT
            cell.alignment = Alignment(horizontal='center', vertical='center')
        for r in rows:
            ws.append(list(r))
        ws.freeze_panes = freeze
        for i, w in enumerate(widths or [], start=1):
            ws.column_dimensions[get_column_letter(i)].width = w
        ws.auto_filter.ref = 'A1:%s%d' % (get_column_letter(len(headers)), len(rows) + 1)
        return ws

    # ① 质量概览
    overview = [
        ('送检数量', q.get('total') or 0, '区间内所有检验单的送检数量合计'),
        ('合格数量', q.get('pass') or 0, '区间内所有检验单的合格数量合计'),
        ('非合格数量', q.get('bad') or 0, '区间内所有检验单的不合格数量合计(注:不等于 送检−合格)'),
        ('合格率(%)', q.get('passRate') or 0, '合格数量 / 送检数量'),
        ('数据区间', '%s ~ %s' % (rng.get('from'), rng.get('to')), '与「我的桌面」默认区间一致(近 30 天)'),
        ('统计粒度', rng.get('bucket'), 'day=按天 / week / month,随区间长度自动切换'),
    ]
    sheet('质量概览', ['指标', '值', '口径'], overview, widths=[16, 22, 56])
    ws = wb['质量概览']
    for r in range(2, ws.max_row + 1):
        ws.cell(row=r, column=1).font = Font(bold=True)
        ws.cell(row=r, column=2).alignment = Alignment(horizontal='right')

    # ② 检验结果分布
    by_result = [(x.get('name'), x.get('value')) for x in (q.get('byResult') or [])]
    sheet('检验结果分布', ['结果', '数量'], by_result, widths=[18, 16])

    # ③ 送检/合格趋势
    trend = [(x.get('date'), x.get('added'), x.get('done'),
              (x.get('done') or 0) - (x.get('added') or 0) if (x.get('added') or 0) else 0)
             for x in (q.get('trend7') or [])]
    ws3 = sheet('送检合格趋势', ['日期', '送检', '合格', '差额(合格−送检)'], trend, widths=[12, 14, 14, 18])
    # 合格率列由公式算(打开就能核对,不写死结果)
    ws3.cell(row=1, column=5, value='合格率').fill = HEAD_FILL
    ws3.cell(row=1, column=5).font = HEAD_FONT
    ws3.cell(row=1, column=5).alignment = Alignment(horizontal='center')
    for r in range(2, ws3.max_row + 1):
        ws3.cell(row=r, column=5, value='=IFERROR(C%d/B%d,"")' % (r, r)).number_format = '0.0%'
    ws3.column_dimensions['E'].width = 10

    # ④ 不合格物料 TOP
    defect = [(i + 1, x.get('name'), x.get('value')) for i, x in enumerate(q.get('defectItems') or [])]
    sheet('不合格物料TOP', ['序号', '物料/供应商名称', '不合格数量'], defect, widths=[8, 30, 14])

    # ⑤ 数据说明
    notes = [
        ('导出内容', '「我的桌面」→ 质量模块:质量概览 / 检验结果分布 / 送检合格趋势 / 不合格物料 TOP'),
        ('数据源', 'MES 后端 GET /api/dashboard/stats?from=%s&to=%s 的 quality 段' % (frm, to)),
        ('与界面关系', '就是「我的桌面」质量卡、以及「质量」页签渲染的同一份数据,未另算一遍'),
        ('导出时间', dt.datetime.now().strftime('%Y-%m-%d %H:%M:%S')),
        ('导出脚本', 'tools/archive/_quality-export-20261018/_export-quality-xlsx.py'),
        ('隐私', '本表**不含人员姓名/联系方式**,只有数量与物料/供应商名称(业务事实数据,按原样保留)'),
        ('注意', '数值为导出时刻的快照;要看最新数据请直接开系统,不要以本表为准做决策'),
    ]
    sheet('数据说明', ['项', '内容'], notes, widths=[16, 92], freeze='A2')

    del wb['Sheet']
    out = os.path.join(outdir, '我的桌面质量数据-' + to.replace('-', '') + '.xlsx')
    wb.save(out)
    print('已写出: ' + out)

    # 回读自检
    wb2 = openpyxl.load_workbook(out, read_only=True, data_only=True)
    fails = []
    if wb2.sheetnames != ['质量概览', '检验结果分布', '送检合格趋势', '不合格物料TOP', '数据说明']:
        fails.append('页签不符: ' + str(wb2.sheetnames))
    ov = {r[0]: r[1] for r in wb2['质量概览'].iter_rows(min_row=2, values_only=True)}
    if str(ov.get('送检数量')) != str(q.get('total')):
        fails.append('送检数量不一致: 表=%s 接口=%s' % (ov.get('送检数量'), q.get('total')))
    if str(ov.get('合格数量')) != str(q.get('pass')):
        fails.append('合格数量不一致: 表=%s 接口=%s' % (ov.get('合格数量'), q.get('pass')))
    if str(ov.get('非合格数量')) != str(q.get('bad')):
        fails.append('非合格数量不一致: 表=%s 接口=%s' % (ov.get('非合格数量'), q.get('bad')))
    n_trend = sum(1 for _ in wb2['送检合格趋势'].iter_rows(min_row=2))
    if n_trend != len(q.get('trend7') or []):
        fails.append('趋势点数不一致: 表=%d 接口=%d' % (n_trend, len(q.get('trend7') or [])))
    sheet_names = wb2.sheetnames
    wb2.close()

    summary = {
        'generatedAt': dt.datetime.now().isoformat(timespec='seconds'),
        'range': rng,
        'quality': {'total': q.get('total'), 'pass': q.get('pass'), 'bad': q.get('bad'), 'passRate': q.get('passRate')},
        'sheets': sheet_names,
        'trendPoints': n_trend,
        'defectItems': len(defect),
        'file': os.path.basename(out),
        'failed': fails,
    }
    with open(os.path.join(outdir, '_export-summary.json'), 'w', encoding='utf-8', newline='') as f:
        json.dump(summary, f, ensure_ascii=False, indent=2)

    if fails:
        for x in fails:
            print('[FAIL] ' + x)
        return 1
    print('[OK] 自检通过: 5 个页签, 概览三项与接口一致, 趋势 %d 点, 不合格项 %d 条' % (n_trend, len(defect)))
    return 0


if __name__ == '__main__':
    sys.exit(main())
