# -*- coding: utf-8 -*-
"""_dump-docx-structure.py — 抽取 8 份检验规范 .docx 的结构骨架(段落标题 + 表格形状)

用途:为「把所有成品检验规范整理成**一个文档格式**」找共同骨架(2026-10-09 用户任务)。
输出 UTF-8 文件,避免 PowerShell 控制台 GBK 乱码。
"""
import glob
import io
import os
from docx import Document

SRC = os.path.join("tools", "archive", "_qcspec", "docx")
OUT = os.path.join("tools", "archive", "_qcspec", "_structure.txt")
MAX_PARA = 45
MAX_CELL = 10


def cell_text(c):
    return " ".join((c.text or "").split())[:40]


def main():
    files = sorted(glob.glob(os.path.join(SRC, "*.docx")))
    with io.open(OUT, "w", encoding="utf-8") as f:
        for path in files:
            d = Document(path)
            f.write(u"\n" + u"=" * 100 + u"\n")
            f.write(u"### %s\n" % os.path.basename(path))
            f.write(u"段落 %d 个 / 表格 %d 张 / 节 %d 个\n" % (len(d.paragraphs), len(d.tables), len(d.sections)))
            f.write(u"---- 段落(非空) ----\n")
            n = 0
            for i, p in enumerate(d.paragraphs):
                t = " ".join((p.text or "").split())
                if not t:
                    continue
                n += 1
                if n > MAX_PARA:
                    f.write(u"  ...(还有更多段落,略)\n")
                    break
                f.write(u"  [%02d] (%s) %s\n" % (i, p.style.name if p.style else "?", t[:80]))
            f.write(u"---- 表格 ----\n")
            for ti, tb in enumerate(d.tables):
                f.write(u"  表%d: %d 行 x %d 列\n" % (ti + 1, len(tb.rows), len(tb.columns)))
                for ri, row in enumerate(tb.rows[:6]):
                    cells = [cell_text(c) for c in row.cells[:MAX_CELL]]
                    f.write(u"      r%d: %s\n" % (ri, u" | ".join(cells)))
                if len(tb.rows) > 6:
                    f.write(u"      ...(共 %d 行)\n" % len(tb.rows))
    print("OK ->", OUT)


if __name__ == "__main__":
    main()
